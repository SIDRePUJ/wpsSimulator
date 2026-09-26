"""Profile downloaded public evidence without exporting individual CNA records."""

from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import statistics
import unicodedata
import zipfile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
RAW = ROOT / "data" / "raw"
REPORTS = ROOT / "reports" / "raw"
MUNICIPALITY = "13442"
KEY_FIELDS = ("P_MUNIC", "UC_UO", "ENCUESTA", "COD_VEREDA")


def number(value):
    if value is None or str(value).strip().upper() in {"", "NA", "NAN"}:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def key(row):
    return tuple(str(row.get(field, "")).strip() for field in KEY_FIELDS)


def normalize(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in value if not unicodedata.combining(char)).casefold().strip()


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def crop_dictionary():
    path = REPORTS / "dane_cna2014_data_dictionary.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Cod_Cultivo"]
    mapping = {
        str(row[1]).strip(): str(row[2]).strip()
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[1] is not None and row[2] is not None
    }
    workbook.close()
    return mapping


def cna_profile():
    path = RAW / "dane_cna2014_bolivar_csv.zip"
    names = crop_dictionary()
    upa = {}
    upa_count = 0
    upa_duplicates = 0
    upa_missing_key = 0
    source_district = Counter()
    drought_difficulty = Counter()
    both_district_and_drought = 0
    with zipfile.ZipFile(path) as archive:
        with io.TextIOWrapper(archive.open("S01_15_Unidad_productora_.csv"), encoding="utf-8", newline="") as stream:
            for row in csv.DictReader(stream):
                if row["P_MUNIC"] != MUNICIPALITY:
                    continue
                upa_count += 1
                k = key(row)
                if any(not part for part in k):
                    upa_missing_key += 1
                if k in upa:
                    upa_duplicates += 1
                upa[k] = {
                    "district": row.get("P_S11P124_SP10") == "1",
                    "drought": row.get("P_S11P126_SP4") == "1",
                }
                if upa[k]["district"] and upa[k]["drought"]:
                    both_district_and_drought += 1
                source_district[row.get("P_S11P124_SP10", "")] += 1
                drought_difficulty[row.get("P_S11P126_SP4", "")] += 1

        crop_count = 0
        crop_keys = set()
        crop_matches = 0
        valid_yield = 0
        crop_stats = defaultdict(lambda: Counter())
        district_yield_upa = defaultdict(set)
        rice_district_yields = []
        rice_district_harvested_ha = 0.0
        rice_district_production_t = 0.0
        with io.TextIOWrapper(archive.open("S06A_Cultivos_.csv"), encoding="utf-8", newline="") as stream:
            for row in csv.DictReader(stream):
                if row["P_MUNIC"] != MUNICIPALITY:
                    continue
                crop_count += 1
                k = key(row)
                crop_keys.add(k)
                match = upa.get(k)
                if match is not None:
                    crop_matches += 1
                code = row.get("P_S6P46", "")
                label = names.get(code, "UNMAPPED: " + code)
                stats = crop_stats[label]
                stats["records"] += 1
                if number(row.get("AREA_COSECHADA")) not in (None, 0):
                    stats["harvested_area_records"] += 1
                if number(row.get("P_S6P59_UNIF")) not in (None, 0):
                    stats["positive_yield_records"] += 1
                    valid_yield += 1
                    if match and match["district"]:
                        stats["positive_yield_records_in_district_source_upa"] += 1
                        district_yield_upa[label].add(k)
                        if label == "Arroz verde":
                            rice_district_yields.append(number(row["P_S6P59_UNIF"]))
                            rice_district_harvested_ha += number(row["AREA_COSECHADA"]) or 0
                            rice_district_production_t += number(row["P_S6P57A"]) or 0
                if match and match["district"]:
                    stats["records_in_district_source_upa"] += 1
                if match and match["drought"]:
                    stats["records_in_drought_difficulty_upa"] += 1

        household_count = 0
        household_keys = set()
        with io.TextIOWrapper(archive.open("S15H_Hogares_.csv"), encoding="utf-8", newline="") as stream:
            for row in csv.DictReader(stream):
                if row["P_MUNIC"] == MUNICIPALITY:
                    household_count += 1
                    household_keys.add(key(row))

    return {
        "municipality_code": MUNICIPALITY,
        "upa_rows": upa_count,
        "upa_unique_composite_keys": len(upa),
        "upa_duplicate_composite_keys": upa_duplicates,
        "upa_missing_composite_key": upa_missing_key,
        "upa_irrigation_district_source_yes": source_district["1"],
        "upa_2013_drought_difficulty_yes": drought_difficulty["1"],
        "upa_both_district_source_and_drought_difficulty": both_district_and_drought,
        "crop_rows": crop_count,
        "crop_unique_composite_keys": len(crop_keys),
        "crop_rows_matching_upa_composite_key": crop_matches,
        "crop_rows_with_positive_yield": valid_yield,
        "household_rows": household_count,
        "household_unique_composite_keys": len(household_keys),
        "household_keys_matching_upa": len(household_keys & upa.keys()),
        "household_keys_in_district_source_upa": sum(1 for household_key in household_keys if upa.get(household_key, {}).get("district")),
        "rice_district_source_upa_with_positive_yield_and_household": len(district_yield_upa["Arroz verde"] & household_keys),
        "rice_district_source_positive_yield_summary": {
            "records": len(rice_district_yields),
            "median_yield_t_ha": statistics.median(rice_district_yields) if rice_district_yields else None,
            "harvested_area_ha_in_cna_records": rice_district_harvested_ha,
            "production_t_in_cna_records": rice_district_production_t,
            "aggregate_yield_t_ha": rice_district_production_t / rice_district_harvested_ha if rice_district_harvested_ha else None,
        },
        "crop_counts": {
            label: dict(stats) | {"unique_upa_with_positive_yield_in_district_source": len(district_yield_upa[label])}
            for label, stats in sorted(crop_stats.items(), key=lambda item: (-item[1]["records"], item[0]))
        },
        "note": "UPA, crop and household tables have different units. A matching composite key does not imply one household per UPA.",
    }


def eva_profile(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["BasePagina"]
    header_row = next(
        index for index, values in enumerate(sheet.iter_rows(values_only=True), 1)
        if values[3] == "Municipio" and values[5] == "Cultivo"
    )
    headers = [str(cell or "").strip() for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    selected = []
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if str(values[2] or "").strip() == MUNICIPALITY:
            row = dict(zip(headers, values))
            selected.append({
                "crop_detail": row[headers[4]],
                "crop": row[headers[5]],
                "cycle": row[headers[6]],
                "year": row[headers[9]],
                "period": row[headers[10]],
                "sown_ha": number(row[headers[11]]),
                "harvested_ha": number(row[headers[12]]),
                "production_t": number(row[headers[13]]),
                "yield_t_ha": number(row[headers[14]]),
            })
    workbook.close()
    years = Counter(str(row["year"]) for row in selected)
    latest_year = max((int(year) for year in years if year.isdigit()), default=None)
    latest = [row for row in selected if str(row["year"]) == str(latest_year)]
    latest_crop_area = Counter()
    for row in latest:
        latest_crop_area[str(row["crop"])] += row["harvested_ha"] or 0
    rice = [row for row in selected if "arroz" in normalize(row["crop"])]
    return {
        "file": path.name,
        "local_rows": len(selected),
        "rows_by_year": dict(sorted(years.items())),
        "latest_year": latest_year,
        "latest_top_crops_by_harvested_ha": latest_crop_area.most_common(12),
        "rice_rows": rice,
        "header_row": header_row,
    }


def ideam_profile():
    path = RAW / "ideam_station_catalog.csv"
    stations = []
    with path.open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            if normalize(row["Municipio"]) == "maria la baja":
                stations.append({field: row[field] for field in (
                    "Codigo", "Nombre", "Categoria", "Estado", "LATITUD", "LONGITUD", "Fecha_instalacion", "Fecha_suspension"
                )})
    active_rain = [row for row in stations if row["Estado"] == "Activa" and "Pluvio" in row["Categoria"]]
    return {"local_station_count": len(stations), "active_rain_station_count": len(active_rain), "active_rain_stations": active_rain,
            "note": "Catalog installation and status do not establish observation completeness or quality."}


def historical_eva_profile():
    path = RAW / "agronet_eva_2007_2018.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["FINAL"]
    local_rows = []
    for values in sheet.iter_rows(min_row=3, values_only=True):
        if str(values[2] or "") == MUNICIPALITY:
            local_rows.append({
                "crop": values[6],
                "system": values[7],
                "year": values[8],
                "period": values[9],
                "sown_ha": number(values[10]),
                "harvested_ha": number(values[11]),
                "production_t": number(values[12]),
                "yield_t_ha": number(values[13]),
            })
    workbook.close()
    rice_2013 = [row for row in local_rows if row["year"] == 2013 and normalize(row["system"]) == "arroz riego"]
    rice_area = sum(row["harvested_ha"] or 0 for row in rice_2013)
    rice_production = sum(row["production_t"] or 0 for row in rice_2013)
    return {
        "local_rows": len(local_rows),
        "rows_by_year": dict(sorted(Counter(str(row["year"]) for row in local_rows).items())),
        "rice_irrigated_2013_rows": rice_2013,
        "rice_irrigated_2013_harvested_ha": rice_area,
        "rice_irrigated_2013_production_t": rice_production,
        "rice_irrigated_2013_aggregate_yield_t_ha": rice_production / rice_area if rice_area else None,
        "note": "Historical EVA reports municipal production by sowing period, whereas CNA reports UPA crop records; comparability is not guaranteed.",
    }


def nasa_profile():
    path = RAW / "nasa_power_maria_la_baja_2013_2024_daily.csv"
    by_year = defaultdict(lambda: {"days": 0, "precipitation_mm": 0.0, "missing_precipitation_days": 0})
    with path.open(encoding="utf-8", newline="") as stream:
        for line in stream:
            if line.strip() == "-END HEADER-":
                break
        for row in csv.DictReader(stream):
            year = row["YEAR"]
            stats = by_year[year]
            stats["days"] += 1
            rain = number(row["PRECTOTCORR"])
            if rain is None or rain == -999:
                stats["missing_precipitation_days"] += 1
            else:
                stats["precipitation_mm"] += rain
    return {"by_year": dict(sorted(by_year.items())),
            "note": "NASA POWER is a gridded MERRA-2/CERES proxy, not an IDEAM station observation or observed district water supply."}


def main():
    files = sorted((*RAW.iterdir(), *REPORTS.iterdir()))
    inventory = [{"path": str(path.relative_to(ROOT)).replace("\\", "/"), "bytes": path.stat().st_size, "sha256": sha256(path)}
                 for path in files if path.is_file()]
    profile = {
        "generated_utc_date": datetime.now(timezone.utc).date().isoformat(),
        "source_files": inventory,
        "cna_2014": cna_profile(),
        "eva_2024_final_base": eva_profile(RAW / "upra_eva_2019_2024.xlsx"),
        "eva_2025_base": eva_profile(RAW / "upra_eva_2019_2025.xlsx"),
        "eva_2007_2018_historical": historical_eva_profile(),
        "ideam_catalog": ideam_profile(),
        "nasa_power": nasa_profile(),
    }
    output = ROOT / "data" / "derived" / "source_profile.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(profile, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(output)
    print("CNA UPA/crops:", profile["cna_2014"]["upa_rows"], profile["cna_2014"]["crop_rows"])
    print("EVA local rows:", profile["eva_2024_final_base"]["local_rows"], profile["eva_2025_base"]["local_rows"])


if __name__ == "__main__":
    main()
